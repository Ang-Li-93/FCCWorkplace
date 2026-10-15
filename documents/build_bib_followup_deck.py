#!/usr/bin/env python3
"""Build the 6 Oct 2026 MAPS/PixESL follow-up deck (2026_10_06_MAPS.pptx) on the BNL template of 2026_08_04_MAPS.pptx.

Content: documents/BIB_LCC_FOLLOWUP_SLIDES_TEXT.md (13 slides + the appendix table as 2 backup slides); plots from
outputs/BIB/plots_LCC_truth/ and fcc_maps_wrapper_pixesl/results/bib_ipc_lcc_truth/event_displays/.
Needs python-pptx (not in the system Python):  python3.12 -m venv v && v/bin/pip install python-pptx && v/bin/python documents/build_bib_followup_deck.py [out.pptx]
Render for checking (no LibreOffice at SDCC): apptainer pull docker://lscr.io/linuxserver/libreoffice:latest, then
apptainer exec --bind $PWD:/work libreoffice_latest.sif bash -c "cd /work && HOME=/tmp soffice --headless --convert-to pdf out.pptx".
"""
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml

WS = Path("/gpfs/mnt/gpfs01/usfcc/ali3/FCCWorkplace")
TEMPLATE = WS / "2026_08_04_MAPS.pptx"
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else WS / "2026_10_06_MAPS.pptx"
P_TRUTH = WS / "outputs/BIB/plots_LCC_truth"
P_BIBS = WS / "outputs/BIB/bibstudies_LCC_labelled"
P_DISP = WS / "fcc_maps_wrapper_pixesl/results/bib_ipc_lcc_truth/event_displays"

TEAL = RGBColor(0x10, 0x5C, 0x78)      # theme accent1
CYAN = RGBColor(0x00, 0xAD, 0xDC)      # accent2
ORANGE = RGBColor(0xF6, 0x8B, 0x1F)    # accent4
MAGENTA = RGBColor(0xB7, 0x24, 0x67)   # accent5
DARK = RGBColor(0x26, 0x26, 0x26)
GREY = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CARD = RGBColor(0xEE, 0xF4, 0xF7)      # light tint of accent1
CARD2 = RGBColor(0xFD, 0xF3, 0xE8)     # light tint of accent4

FOOTER = "IPC @ Z · LCC V106.2 · central mdi_v2 simulation · 19,673 BX · SimTrackerHits, 1 keV simulation cut"
A_NS = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'

# ----------------------------------------------------------------------------- helpers

def rgb(run, color):
    run.font.color.rgb = color


def add_runs(p, runs, size, color=DARK, bold=False, italic=False):
    """runs: str or list of (text, style-dict) with keys b, i, c (RGBColor), link (url), sz."""
    if isinstance(runs, str):
        runs = [(runs, {})]
    for item in runs:
        text, st = (item, {}) if isinstance(item, str) else item
        r = p.add_run()
        r.text = text
        r.font.size = Pt(st.get("sz", size))
        r.font.bold = st.get("b", bold)
        r.font.italic = st.get("i", italic)
        if st.get("c", color) is not None:
            rgb(r, st.get("c", color))
        if st.get("link"):
            r.hyperlink.address = st["link"]
    return p


def set_bullet(p, level=0):
    """Template bullets: Wingdings 'v' (diamond) at level 1, 'Ø' (arrow) at level 2."""
    pPr = p._p.get_or_add_pPr()
    marL, indent = (342900, 342900) if level == 0 else (742950, 285750)
    pPr.set("marL", str(marL))
    pPr.set("indent", str(-indent))
    if level:
        pPr.set("lvl", str(level))
    char = "v" if level == 0 else "Ø"
    pPr.append(parse_xml(f'<a:buFont {A_NS} typeface="Wingdings" pitchFamily="2" charset="2"/>'))
    pPr.append(parse_xml(f'<a:buChar {A_NS} char="{char}"/>'))


def no_bullet(p):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", "0"); pPr.set("indent", "0")
    pPr.append(parse_xml(f'<a:buNone {A_NS}/>'))


def textbox(slide, x, y, w, h, name="TextBox", anchor=MSO_ANCHOR.TOP, margin=0.05):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    return tb


def bullets(slide, x, y, w, h, items, size=15, name="Body Bullets", space=7, color=DARK):
    """items: list of (level, runs) or runs (level 0)."""
    tb = textbox(slide, x, y, w, h, name=name)
    tf = tb.text_frame
    first = True
    for it in items:
        level, runs = it if isinstance(it, tuple) and isinstance(it[0], int) else (0, it)
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_before = Pt(space if level == 0 else max(2, space - 4))
        p.line_spacing = 1.0
        set_bullet(p, level)
        add_runs(p, runs, size - (1 if level else 0), color=color)
    return tb


def paragraphs(slide, x, y, w, h, items, size=14, name="Text", space=4, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, color=DARK):
    tb = textbox(slide, x, y, w, h, name=name, anchor=anchor)
    tf = tb.text_frame
    for i, runs in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(space)
        p.alignment = align
        no_bullet(p)
        add_runs(p, runs, size, color=color)
    return tb


def card(slide, x, y, w, h, fill=CARD, name="Card"):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.name = name
    sp.adjustments[0] = 0.05
    st = sp._element.find(qn("p:style"))
    if st is not None:
        sp._element.remove(st)
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    sp.line.fill.background()
    sp.shadow.inherit = False
    tf = sp.text_frame
    tf.text = ""
    return sp


def card_text(slide, x, y, w, h, header, body, fill=CARD, hsize=15, bsize=13, name="Card", body_bullets=False, hcolor=TEAL):
    """Card with a bold header line and body paragraphs (list of runs)."""
    card(slide, x, y, w, h, fill=fill, name=name)
    pad = 0.16
    tb = textbox(slide, x + pad, y + pad * 0.6, w - 2 * pad, h - pad, name=name + " text")
    tf = tb.text_frame
    p = tf.paragraphs[0]
    no_bullet(p)
    add_runs(p, header, hsize, color=hcolor, bold=True)
    p.space_after = Pt(3)
    for runs in body:
        q = tf.add_paragraph()
        q.space_before = Pt(4)
        if body_bullets:
            set_bullet(q, 0)
            q._p.get_or_add_pPr().set("marL", "228600"); q._p.get_or_add_pPr().set("indent", "-228600")
        else:
            no_bullet(q)
        add_runs(q, runs, bsize)
    return tb


def picture(slide, path, x, y, w=None, h=None, name=None):
    kw = {}
    if w: kw["width"] = Inches(w)
    if h: kw["height"] = Inches(h)
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), **kw)
    pic.name = name or Path(path).stem
    return pic


def caption(slide, x, y, w, text, size=10, h=0.3, align=PP_ALIGN.LEFT):
    return paragraphs(slide, x, y, w, h, [text], size=size, name="Caption", space=0, color=GREY, align=align)


def table(slide, x, y, w, col_w, rows, header, size=12, hsize=None, row_h=0.32, margin=0.06, name="Table", align_cols=None, bold_col0=False, hrow_h=None):
    """rows: list of lists of str/runs. col_w: fractions summing to 1."""
    n = len(rows) + 1
    shp = slide.shapes.add_table(n, len(header), Inches(x), Inches(y), Inches(w), Inches(row_h * n))
    shp.name = name
    tbl = shp.table
    for i, f in enumerate(col_w):
        tbl.columns[i].width = Inches(w * f)
    for r in range(n):
        tbl.rows[r].height = Inches(row_h)
    tbl.rows[0].height = Inches(hrow_h or row_h)
    hsize = hsize or size
    for c, txt in enumerate(header):
        cell = tbl.cell(0, c)
        cell.fill.solid(); cell.fill.fore_color.rgb = TEAL
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        add_runs(p, txt, hsize, color=WHITE, bold=True)
        for m in ("margin_left", "margin_right"): setattr(cell, m, Inches(margin))
        cell.margin_top = cell.margin_bottom = Inches(0.03)
    for r, row in enumerate(rows, start=1):
        for c, txt in enumerate(row):
            cell = tbl.cell(r, c)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.text_frame.word_wrap = True
            p = cell.text_frame.paragraphs[0]
            if align_cols and c in align_cols:
                p.alignment = align_cols[c]
            add_runs(p, txt, size, bold=(bold_col0 and c == 0))
            for m in ("margin_left", "margin_right"): setattr(cell, m, Inches(margin))
            cell.margin_top = cell.margin_bottom = Inches(0.03)
    return shp


def stat(slide, x, y, w, h, number, label, color=TEAL, nsize=36, lsize=11):
    card(slide, x, y, w, h, name="Stat " + number)
    tb = textbox(slide, x, y + 0.05, w, h - 0.1, name="Stat text " + number, anchor=MSO_ANCHOR.MIDDLE)
    tf = tb.text_frame
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; no_bullet(p)
    add_runs(p, number, nsize, color=color, bold=True)
    q = tf.add_paragraph(); q.alignment = PP_ALIGN.CENTER; no_bullet(q); q.space_before = Pt(2)
    add_runs(q, label, lsize, color=GREY)
    return tb


# ----------------------------------------------------------------------------- deck

prs = Presentation(str(TEMPLATE))
layouts = {l.name: l for l in prs.slide_layouts}
L_TITLE = layouts["1_Title Slide_Print-friendly"]
L_ONLY = layouts["Title Only"]

# keep a copy of the template's slide-number placeholder, then drop all template slides
sldnum_xml = None
for shp in prs.slides[1].shapes:
    if shp.is_placeholder and shp.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
        sldnum_xml = copy.deepcopy(shp._element)
assert sldnum_xml is not None
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.rId)
    sldIdLst.remove(sldId)

SLIDES = []


def new_slide(title, notes=None, layout=None):
    s = prs.slides.add_slide(layout or L_ONLY)
    if title is not None:
        t = s.shapes.title
        t.left, t.top, t.width, t.height = Inches(0.55), Inches(0.28), Inches(12.25), Inches(0.82)
        tf = t.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.text = ""
        add_runs(tf.paragraphs[0], title, 28, bold=True)
    sn = copy.deepcopy(sldnum_xml)
    sn.find(qn("p:nvSpPr")).find(qn("p:cNvPr")).set("id", "900")
    s.shapes._spTree.append(sn)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    SLIDES.append(s)
    return s


def footer(s):
    paragraphs(s, 0.62, 7.07, 11.9, 0.3, [FOOTER], size=10, name="Sample footer", space=0, color=GREY)


# ---- 1 Title
s = prs.slides.add_slide(L_TITLE)
SLIDES.append(s)
ph = {p.placeholder_format.idx: p for p in s.placeholders}
ph[0].text_frame.text = ""
add_runs(ph[0].text_frame.paragraphs[0], "Beam-induced background for the MAPS vertex demonstrator", 48, color=None, bold=True)
ph[10].text_frame.text = ""
add_runs(ph[10].text_frame.paragraphs[0],
         "Answers to the 4 Aug comments: what the axes mean · why layer 4 is above layer 3 · ROOT files for event displays", 18)
ph[1].text_frame.text = ""
add_runs(ph[1].text_frame.paragraphs[0], "Ang Li, Chris Rasmussen, Punit Sharma — LDRD WP1, 06 Oct. 2026", 16)
sn = copy.deepcopy(sldnum_xml); sn.find(qn("p:nvSpPr")).find(qn("p:cNvPr")).set("id", "900"); s.shapes._spTree.append(sn)
s.notes_slide.notes_text_frame.text = (
    "Follow-up to the 4 Aug 2026 review. Sample: IPC at the Z pole, LCC V106.2 lattice, central mdi_v2 "
    "simulation, all 19,673 bunch crossings. All numbers from outputs/BIB/plots_LCC_truth/summary_numbers.md; "
    "errors are BX-to-BX (bursts included).")

# ---- 2 Two comments and short answers
s = new_slide("The two comments and the short answers",
              notes='"Primaries" = the e+/e- pairs from GuineaPig, about 1,120 per bunch crossing, median pT 1 MeV.')
bullets(s, 0.62, 1.2, 12.1, 5.75, [
    [("\u201cDo all the Y axes make sense?\u201d", {"b": True})],
    (1, ["Of the nine plots shown, two are right as labelled and three only lack labels or units. The two \u201coccupancy\u201d plots are not occupancies and the two hit maps are not densities."]),
    (1, ["Every axis is now defined and checked against an independent recount; 54 of the 69 histograms are redrawn with correct labels."]),
    [("\u201cWhy is layer 4 higher than layer 3 \u2014 low-pT loopers?\u201d", {"b": True})],
    (1, ["It is real, but it is not loopers. 94% of the layer-4 hits are made by sub-MeV secondary electrons, not by particles from the beam line; about half are tagged directly as photons absorbed in the silicon."]),
    (1, ["Per cm\u00b2 they are only 39% rarer on layer 4 than on layer 3, and layer 4 has 4.4 times more silicon."]),
    (1, ["The calculation from field and geometry: a particle needs pT > 39 MeV to reach layer 3 and > 95 MeV to reach layer 4. Per crossing 0.024 and 0.0036 primaries do, and the simulation confirms them one by one."]),
    (1, ["With a pT cut of 1 MeV on the track that makes the hit: layer 3 has 0.103 and layer 4 has 0.028 hits per crossing \u2014 the expected order."]),
    [("ROOT files", {"b": True}), " with all hits and their truth, plus an event-display script, are public: https://angli-share.web.cern.ch/FCC/MAPS/BIB_IPC_LCC/truth_ntuple/"],
], size=16)
footer(s)

# ---- 3 Y axes table
s = new_slide("Y axes of the plots shown on 4 Aug", notes=(
    "The missing labels come from how the plots were drawn: drawhits.py sets readable axis titles only in its own "
    "draw step, and our merged redraw used the titles stored in the histograms. All 69 histograms are audited in the "
    "backup slides; 54 are redrawn with correct titles in bibstudies_LCC_labelled/; 15 are not redrawn: 9 empty, 4 "
    "broken by a module-axis offset, 2 uninformative (listed in SKIPPED.md)."))
rows = [
    ["5 left", "h_avg_hits_x_layer", "avg hits/evt", "SimTrackerHits per bunch crossing, summed over the whole layer", "correct; a layer total, not a density"],
    ["5 right", "h_avg_occ_x_layer", "avg occupancy/evt [%]", "100 × fired pixels ÷ number of sensors", "not an occupancy"],
    ["7 left", "hit_xy", "hits/(11×11) mm² per event", "hits per crossing in an 11.28 mm map cell of the end-on projection, summed along z", "not a density"],
    ["7 right", "hit_zr", "hits/(8×11) mm² per event", "hits per crossing in an 8.64 × 11.28 mm cell; “r” is r × sign(x), one band = half a layer", "not a density"],
    ["8 left", "h_hit_particle_ID", "(none)", "hits per crossing by species of the stored MC particle the hit is linked to", "label missing; names the ancestor"],
    ["8 right", "h_hit_E_keV", "hits", "hits per crossing per keV, five layers summed", "unit missing; describes L0–L2"],
    ["9 left", "hist_hit_t", "(none)", "hits per crossing per 0.1 ns, five layers summed, 0–20 ns", "label missing; hides the late MAPS hits"],
    ["9 right", "h_occ_tot", "(none)", "fraction of crossings per bin of 100 × fired pixels ÷ 4,870 sensors", "not an occupancy"],
    ["10", "h_avg_firing_cells_x_layer", "avg cellsFired/evt", "distinct fired pixels per crossing (25 µm inner, 50×150 µm outer)", "correct; the table beside it was not"],
]
table(s, 0.62, 1.2, 12.1, [0.09, 0.17, 0.16, 0.37, 0.21], rows,
      ["4 Aug slide", "plot", "axis as shown", "what it is", "verdict"], size=11, row_h=0.36, name="Axis table")
bullets(s, 0.62, 5.45, 12.1, 1.5, [
    ["“evt” always means one bunch crossing. A “hit” is one Geant4 track in one sensor, not a particle and not a pixel."],
    ["Every filled histogram closes bin by bin against an independent recount from the truth ntuple, which matches the raw files hit by hit."],
], size=14)
footer(s)

# ---- 4 Occupancy
s = new_slide("“Occupancy [%]” was fired pixels per sensor", notes=(
    'The "27%" would even exceed 100% for one crossing (95 pixels in 90 sensors). These occupancies count one pixel '
    "per hit: no charge sharing or cluster size is simulated."))
picture(s, P_TRUTH / "occupancy_per_layer.png", 0.55, 1.3, w=6.6)
bullets(s, 7.3, 1.2, 5.5, 5.7, [
    ["bib-studies counts distinct cellIDs (pixels in this sample) and divides by the ", ("number of sensors", {"b": True}),
     ". Layer 0: 24.0 fired pixels ÷ 90 sensors = 26.7 “%”."],
    ["True pixel occupancy per crossing (fired pixels ÷ all pixels): ", ("8.1e-7", {"b": True}), " in layer 0 (25 µm pixels), ",
     ("6.9e-9", {"b": True}), " in layer 3 and ", ("2.6e-9", {"b": True}), " in layer 4 (the 50×150 µm pixels of the sample)."],
    ["At the 20 µm MAPS pitch: 3.8e-10 (layer 3) and 1.5e-10 (layer 4)."],
    ["Sensors with at least one hit per crossing: 27.7 of 4,870 (0.57%), not “0.8% of modules”. Layer 3: 0.23 of 736. Layer 4: 0.41 of 3,264."],
    ["85% of crossings leave layer 3 empty and 68% leave layer 4 empty."],
], size=15)
caption(s, 0.62, 4.85, 6.5, "Fired pixels per crossing divided by the number of pixels of each layer, at the sample’s pitch and at 20 µm.", size=10, h=0.45)
footer(s)

# ---- 5 Corrections
s = new_slide("Corrections to statements on the 4 Aug slides", notes=(
    "Like-for-like table for the validation slide, SimTrackerHits per crossing, L0...L4: same 2,129 BX and 1.8 keV cut "
    "as A. Ilg (his chart has 78,506 entries, which we reproduce exactly); our values for that subset: 24.00 / 7.55 / "
    "4.65 / 0.279 +- 0.042 / 0.390 +- 0.017. This work, all 19,673 BX, 1.8 keV: 24.05 / 7.73 / 4.82 / 0.246 / 0.404. "
    "Without the threshold: 24.32 / 7.91 / 4.98 / 0.306 / 0.565. The old 'this work' column (24.0 / 7.8 / 4.9 / 0.31 / "
    "0.57) mixed fired cells and hits. Layer 4 is above layer 3 in his subset too (0.39 vs 0.28). Hit rates on sensitive "
    "silicon: 1.13e-4 (L3) and 4.70e-5 (L4) hits/cm2/BX = 4.2 and 1.7 kHz/cm2 at 37.04 MHz (the crossing rate assumed "
    "for LCC V106 in A. Ilg's talk)."))
rows = [
    [[("“Raw SIM hits, no threshold”", {"i": True})], "The central simulation drops every Geant4 step below 1 keV. No hit is below 1 keV; there is no sub-keV tail in the CSV."],
    [[("“MIP-like Landau, peak ~15 keV”", {"i": True})], "True for layers 0–2. On the MAPS layers 20% (L3) and 29% (L4) of the hits deposit 1–2 keV."],
    [[("“~99% of hits are e⁺/e⁻”", {"i": True})], "That is the species of the stored ancestor. On layer 4 only 6% of the hits are made by that particle itself."],
    [[("“Timing is prompt, within 1–2 ns”", {"i": True})], "True for the five-layer sum. Median hit time is 4.4 ns on layer 3 and 7.0 ns on layer 4; 9% and 13% arrive after the next crossing (27 ns)."],
    [[("“Per cm² L3 is ~2.6× busier; L4 has ~4.8× more silicon”", {"i": True})], "2.40× and 4.43× with the sensitive area (2.6 and 4.8 are for bare cylinders)."],
    [[("“Validated against the official analysis, independent analyzer”", {"i": True})], "The reference chart is the same events and the same script with a 1.8 keV threshold; we reproduce it exactly (78,506 entries in the first 2,129 crossings)."],
]
table(s, 0.62, 1.25, 12.1, [0.33, 0.67], rows, ["Said on 4 Aug", "What the data say"], size=13, row_h=0.62, hrow_h=0.4, name="Corrections table")
bullets(s, 0.62, 5.7, 12.1, 1.3, [
    ["Like-for-like check of the validation chart (first 2,129 crossings, 1.8 keV threshold, hits per crossing L0 \u2026 L4): A. Ilg has 78,506 entries, "
     "reproduced exactly: 24.00 / 7.55 / 4.65 / 0.279 / 0.390. All 19,673 crossings at 1.8 keV: 24.05 / 7.73 / 4.82 / 0.246 / 0.404; without threshold: 24.32 / 7.91 / 4.98 / 0.306 / 0.565."],
], size=13)
footer(s)

# ---- 6 The calculation
s = new_slide("The calculation: which primaries can reach each layer", notes=(
    "Counting also the return arc and later turns of loopers in vacuum, the calculation still puts layer 4 three to seven "
    "times below layer 3. Most primaries with high pT are very forward and leave along z: 41 per crossing have enough pT "
    "for layer 3, but only 0.06% of them cross it inside its length. The helix code was checked against a brute-force "
    "stepper and re-implemented independently (identical counts in all 19,673 BX). The field in the barrel is 2 T to "
    "within 1% (solenoid plus anti-solenoid map), and so are the thresholds."))
bullets(s, 0.62, 1.2, 6.2, 5.7, [
    ["In a 2 T field a particle from the beam line reaches radius r only if ", ("pT [GeV] > 0.15 × B [T] × r [m]", {"b": True}),
     ": 39 MeV for layer 3 (131 mm), 95 MeV for layer 4 (316 mm). Below about 600 MeV it cannot leave the tracker radially: it loops unless it leaves along z."],
    ["Each of the 22 million primaries was propagated as a helix from where GuineaPig hands it over. Reaching the layer inside its length, per crossing: ",
     ("0.024", {"b": True}), " (layer 3) and ", ("0.0036", {"b": True}), " (layer 4)."],
    ["The simulation agrees particle by particle: 91% (L3) and 86% (L4) of the predicted primaries leave a hit there; the misses are explained by sensor gaps and energy loss."],
    ["So field and geometry put layer 4 about ", ("seven times below", {"b": True}), " layer 3. The simulation has it ", ("1.85 times above", {"b": True}), "."],
], size=15)
picture(s, P_TRUTH / "helix_prediction_vs_sim.png", 7.0, 1.2, w=5.8)
caption(s, 7.0, 4.85, 5.8, "Primaries per crossing predicted by the helix calculation to cross each layer, and found in the simulation.", size=10, h=0.45)
footer(s)

# ---- 7 What makes the hits
s = new_slide("What makes the hits", notes=(
    '"Sub-MeV secondary" = a track created in the simulation below 1 MeV; these are not stored in the truth record, the '
    "hit is flagged, and its momentum is known from the hit itself. Blue = the GuineaPig primary itself, orange = a "
    "stored secondary (> 1 MeV)."))
picture(s, P_TRUTH / "hits_per_layer_by_origin.png", 0.55, 1.2, w=6.0)
picture(s, P_TRUTH / "pt_spectrum_L3_L4.png", 6.75, 1.2, w=6.05)
bullets(s, 0.62, 4.35, 12.1, 2.6, [
    ["Layer 3: ", ("0.306", {"b": True}), " hits per crossing = 0.108 from charged particles that are stored in the truth + 0.198 from sub-MeV secondary electrons (", ("65%", {"b": True}), ")."],
    ["Layer 4: ", ("0.565", {"b": True}), " = 0.033 + 0.532 (", ("94%", {"b": True}), ")."],
    ["Hits made by the stored charged particles follow the expected order: three times fewer on layer 4 (six times fewer for the primaries themselves). The whole excess is in the sub-MeV electrons."],
    ["82% of the layer-3 hits and 97% of the layer-4 hits are made by tracks below the pT needed to reach the layer."],
], size=15)
footer(s)

# ---- 8 pT cut
s = new_slide("The pT cut", notes=(
    "The stored hit momentum is the mean of entry and exit momentum, so for an electron that stops in the sensor it is "
    "half the entry value; this only matters for cuts below 1 MeV. A cut at the generator level on the pT of the primary "
    "behind the hit gives 1.76 (> 1 MeV), 1.64 (> 5 MeV), 0.87 (> 10 MeV), 0.31 (> 20 MeV): the layer-4 hits do descend "
    "from low-pT pairs (median 5 MeV), but those pairs never reach layer 4; they act through the radiation they produce "
    "where they hit material. For a common cut above 95 MeV layer 4 is higher again (1.5 +- 0.1): the few loopers that do "
    "reach it cross the twice longer layer more often (0.018 vs 0.012 hits per crossing)."))
rows = [
    ["none", "0.306", "0.565", "1.85 ± 0.06"],
    ["> 1 MeV", "0.103", "0.028", "0.27 ± 0.05"],
    ["> 10 MeV", "0.087", "0.022", "0.26 ± 0.05"],
    ["> the layer’s own reach threshold (39 / 95 MeV)", "0.055", "0.018", "0.34 ± 0.05"],
]
table(s, 0.62, 1.3, 6.1, [0.49, 0.17, 0.17, 0.17], rows,
      ["cut on the pT of the track at the hit", "layer 3", "layer 4", "L4 / L3"], size=13, row_h=0.42, name="pT cut table",
      align_cols={1: PP_ALIGN.CENTER, 2: PP_ALIGN.CENTER, 3: PP_ALIGN.CENTER})
caption(s, 0.62, 3.5, 6.1, "SimTrackerHits per crossing.", size=10)
bullets(s, 0.62, 3.95, 6.1, 3.0, [
    ["A pT cut does restore the expected order. It works because it removes the ", ("keV electrons", {"b": True}), ", not because it removes loopers."],
    ["It is a truth-level cut. A charge threshold does not do the same: above 1,000 e⁻ (3.6 keV) the ratio is still ", ("1.61", {"b": True}), "."],
    ["A cut on the pT of the ", ("primary behind the hit", {"b": True}), " works only above 10 MeV (ratio 1.76 at > 1 MeV, 0.87 at > 10 MeV, 0.31 at > 20 MeV): "
     "the layer-4 hits descend from low-pT pairs that never reach layer 4; they act through the radiation they produce where they hit material."],
], size=15)
picture(s, P_TRUTH / "hits_vs_ptcut.png", 7.0, 1.25, w=5.8)
caption(s, 7.0, 4.9, 5.8, "Hits per crossing on layers 3 and 4 as a function of the cut on the pT of the track at the hit.", size=10, h=0.45)
footer(s)

# ---- 9 What does and does not change the order
s = new_slide("What does and does not change the order", notes=(
    "L4 - L3 = 0.259 +- 0.012 hits per crossing, 21 standard deviations above zero, and not carried by a few crossings: "
    "removing the 20 busiest crossings raises L4/L3 to 2.0, because the bursts (loopers) sit mainly on layer 3. A very "
    "early window also reverses the order (hits before 2 ns: 0.115 vs 0.065), partly by geometry: light needs 1.0 ns to "
    "reach layer 4. Relative to the light time of flight, layer 3 is higher only for windows up to about 1.5 ns; the "
    "layer-4 excess arrives later."))
picture(s, P_TRUTH / "ratio_summary.png", 0.55, 1.2, w=7.2)
bullets(s, 7.95, 1.2, 4.9, 5.7, [
    ["Layer 4 stays above layer 3 however hits are counted: fired 20 µm pixels (0.445 vs 0.259), with a charge threshold, within one bunch spacing."],
    ["Limiting layer 4 to the length of layer 3 brings it just below (0.27 vs 0.31): both the larger circumference and the doubled length are needed."],
    ["Only truth-level selections of charged particles (stored particle, or track pT > 1 MeV) bring the ratio down to the ", ("0.15–0.3", {"b": True}), " expected for particles from the beam line."],
], size=15)
caption(s, 0.62, 5.35, 7.1, "L4 / L3 ratio of hits per crossing for different ways of counting; error bars are BX-to-BX (bootstrap).", size=10, h=0.45)
footer(s)

# ---- 10 Photons in the silicon
s = new_slide("Photons in the silicon make the sub-MeV electrons", notes=(
    "Where the primaries behind these hits end (the photons are often emitted upstream of that point, or arrive after a "
    "detour): forward pairs (median 130 MeV, 33 mrad) follow the field along the beam pipe and shower in it: conical beam "
    "pipe at |z| = 0.28-1.16 m (27% of the layer-4 soft hits), front of the compensating-solenoid cryostat beyond 1.39 m "
    "(19%), the tip of the copper crotch where the beam pipes separate at 1.32 m (18%), copper pipes behind and before "
    "it (20%), LumiCal (5%). About 30% of the layer-4 soft hits arrive later than 14 ns; for the 18-24 ns peak (13%) "
    "timing and direction show radiation returning from the outer detector (the outer barrel at R ~ 2 m or the "
    "endcap-calorimeter face at |z| ~ 3 m; the truth record cannot separate them). The photons themselves are not in the "
    "truth record (below 1 MeV nothing is stored), so the photon origin is an inference from the Auger pairs, the uniform "
    "depth in the silicon, the isolation from any charged track and the timing. This component depends on simulation "
    "choices: 50 um sensitive thickness (it is a volume effect), Auger emission on, the 1 keV cut, and the CAD model of "
    "the beam pipe."))
picture(s, P_TRUTH / "density_vs_radius.png", 0.55, 1.2, h=3.2)
picture(s, P_TRUTH / "edep_L3_L4.png", 5.95, 1.2, h=3.2)
bullets(s, 0.62, 4.5, 12.1, 2.5, [
    ["About two thirds start and stop inside the 50 µm sensor: 29% are 1.6 keV silicon ", ("Auger electrons", {"b": True}),
     " and 35–39% are electrons of a few tens of keV (median 17 keV) whose deposited energy equals their kinetic energy. The rest cross a sensor face."],
    ["About half are tagged directly as photon absorptions: a photo-electron plus the Auger electron at the same point. The Auger electrons are 19% (L3) and 27% (L4) of all “hits”."],
    ["From layer 3 to layer 4 their density per cm² falls to ", ("0.61", {"b": True}), " of its value; for charged-particle hits it falls to ", ("0.07", {"b": True}),
     ". On the silicon wrapper at 2 m it is still half the layer-4 value."],
    ["Hits ∝ photon field × silicon: layer 4 has 2.2 times the silicon per unit length and twice the length of layer 3."],
], size=14)
footer(s)

# ---- 11 Loopers / event displays
s = new_slide("Loopers exist, but they sit mostly on layer 3", notes=(
    "A looper that reaches layer 4 does cross it more often than layer 3 because layer 4 is twice as long (265 vs 182 hits "
    "for the 70 primaries that reach it). That is the reviewers' mechanism; it is worth +0.004 hits per crossing, 1.6% of "
    "the difference. Displays: bx 217846 = busiest crossing (whole tracker), bx 100003 = ordinary crossing with one "
    "isolated soft hit on layer 4 (58% of crossings have no MAPS hit at all), bx 566681 = rare (0.3% of crossings): one "
    "127 MeV primary crossing layers 3 and 4."))
bullets(s, 0.62, 1.2, 5.0, 3.3, [
    ["Second and later crossings of a looping particle are ", ("23%", {"b": True}), " of the layer-3 hits and ", ("3.6%", {"b": True}), " of the layer-4 hits."],
    ["They make the bursts: the busiest 1% of crossings hold 30% of the hits on layer 3 but 10% on layer 4."],
    ["The busiest crossing (112 MAPS hits) is one 477 MeV positron with almost no pz: 75 hits on layer 4 and 24 on layer 3 over 1.4 µs, and 4,000 drift-chamber hits."],
], size=15)
picture(s, P_DISP / "bib_display_bx217846_full.png", 5.75, 1.15, w=7.1)
caption(s, 5.75, 4.05, 7.1, "Busiest crossing (bx 217846), whole tracker: one 477 MeV positron spiralling in the drift chamber, 1.4 µs of hits.", size=10)
picture(s, P_DISP / "bib_display_bx100003.png", 0.62, 4.4, w=5.9)
picture(s, P_DISP / "bib_display_bx566681.png", 6.95, 4.4, w=5.9)
caption(s, 0.62, 6.72, 5.9, "Ordinary crossing (bx 100003): one isolated soft hit on layer 4; 58% of crossings have none.", size=9, h=0.3)
caption(s, 6.95, 6.72, 5.9, "Rare (0.3% of crossings, bx 566681): one 127 MeV primary crossing layers 3 and 4.", size=9, h=0.3)
footer(s)

# ---- 12 ROOT files
URL = "https://angli-share.web.cern.ch/FCC/MAPS/BIB_IPC_LCC/truth_ntuple/"
s = new_slide("ROOT files for event displays", notes=(
    "The barcode column is the truth label asked for AI training (unique per stored particle). For track-finding truth "
    "use hits with by_secondary == 0. On layers 3 and 4 the other hits are isolated (97-99.8% have no charged-track hit "
    "within 1 mm); on layers 0-2 a quarter to a half of them accompany a track. The displays show all hits of a simulated "
    "crossing whatever their time; --tmax 27 keeps one bunch spacing. Plain ROOT works too: "
    'hits->Draw("y:x","bx==217846 && det==0").'))
bullets(s, 0.62, 1.2, 12.1, 3.3, [
    [(URL, {"link": URL}), "  (README, files, script, examples)"],
    [("ipc_LCC_V106p2_central_truth_vertex.root", {"b": True}), " (117 MB): every vertex hit of all 19,673 crossings with position, time, deposited energy, momentum "
     "(and the 20 \u00b5m pixel address on the MAPS layers), and the truth: particle barcode, species, origin, primary ancestor."],
    [("ipc_LCC_V106p2_central_truth.root", {"b": True}), " (574 MB): the same for the whole tracker (Si wrapper, drift chamber). Trees hits, particles, bx; barcode = bx \u00d7 10\u2076 + MC index."],
    [("bib_event_display.py", {"b": True}), ": needs only uproot and matplotlib.  --bx 217846 --tracks   --tmax 27   --busiest 5   --full"],
    ["12 original EDM4hep files (the 6 crossings with most MAPS hits and 6 ordinary ones) for Key4hep displays; the full set is on EOS."],
    ["Plain ROOT works too:  hits->Draw(\"y:x\", \"bx==217846 && det==0\")"],
], size=14)
picture(s, P_DISP / "bib_display_bx217846.png", 3.45, 4.0, w=6.4)
caption(s, 3.45, 6.65, 6.4, "Output of bib_event_display.py for crossing 217846 (vertex detector only).", size=10, align=PP_ALIGN.CENTER)
footer(s)

# ---- 13 Implications and open points
s = new_slide("What this means for the MAPS study, and open points", notes=(
    "For a full-truth re-simulation --part.keepAllParticles alone does nothing with the default particle handler; it "
    'needs SIM.part.userParticleHandler = "" (or minimalKineticEnergy = 0), and SIM.filter.tracker = "edep0". The exact '
    "central ddsim configuration is stored in the runs tree of every central file."))
bullets(s, 0.62, 1.2, 12.1, 5.75, [
    ["Readout load on the MAPS layers: 0.26 (L3) and 0.44 (L4) fired 20 \u00b5m pixels per crossing, mostly isolated hits of several thousand electrons. "
     "A charge threshold does not remove them; being isolated, they should be rejected by pattern recognition (not tested here)."],
    ["The inner layers are different: 85\u201391% of their hits are real charged-particle crossings."],
    ["The soft-photon component must be checked against thickness, Auger setting and the beam-pipe model before it is used for design numbers."],
    ["Caveats of the central sample to keep in mind: 1 keV per-step cut (charge 10\u201320% low below 100 MeV), merged re-entry hits, truth stored only above 1 MeV, "
     "828 of 20,501 crossings missing, no charge sharing between pixels."],
    [("Next", {"b": True})],
    (1, ["re-simulate a subset with full truth and no step cut"]),
    (1, ["SR background"]),
    (1, ["PixESL on the stream with the late hits moved to their crossing"]),
], size=16)
footer(s)

# ---- Backup: axis audit (2 slides)
AUDIT = [
    ["h_avg_hits_x_layer", "hits per crossing, whole layer. Stored error bars are 1.6–3.2× too small (hits come in bursts). The sixth bin is not a layer", "correct"],
    ["h_avg_firing_cells_x_layer", "distinct fired pixels per crossing at the sample’s segmentation (L3 0.249, L4 0.418; at 20 µm 0.259, 0.445)", "correct"],
    ["h_avg_occ_x_layer", "100 × fired pixels ÷ number of sensors", "not an occupancy"],
    ["h_occ_x_layer0–4, h_occ_tot", "fraction of crossings per log-spaced bin of the same ratio, drawn on a linear axis", "not an occupancy"],
    ["h_avg_hits_x_layer0_x_module", "hits per crossing per module", "correct"],
    ["h_avg_hits_x_layer1–4_x_module", "module axis shifted: 38% (L1), 61% (L2), 100% (L3, L4) of the hits in the underflow", "broken"],
    ["h_tot_pu, h_avg_pu_x_layer, h_pu_x_layer0–4", "pile-up was not computed", "empty"],
    ["hit_xy, hist_xy_vs_layer0–4", "hits per crossing per 11.28 mm square of the end-on projection, summed along z. Read as a density it is 13–43× too high", "not a density"],
    ["hit_zr", "hits per crossing per 8.64 × 11.28 mm cell; r carries the sign of x; one cell = half a ring. Compresses the radial fall-off (L0/L3 looks 126:1, is 1,169:1)", "not a density"],
    ["hit_zphi", "hits per crossing per (8.64 mm × 0.0707 rad) cell, five layers on top of each other; stored axis titles swapped", "mislabelled"],
    ["hist_zphi_vs_layer0–4", "the same for one layer; proportional to a real density (one cell = 0.81 cm² on L3, 1.79 cm² on L4). Too finely binned for L3/L4 (< 2 hits per cell)", "labels missing"],
    ["h_hit_E_keV", "hits per crossing per keV, five layers. Empty below 1 keV (simulation cut); line at 1.5–1.75 keV = silicon Auger electrons", "unit missing"],
    ["h_hit_E_MeV, h_hit_E_MeV_map", "100 keV bins: 90% of the hits in the first bin", "uninformative"],
    ["h_hit_E_eV, h_hit_E_eV_map", "covers 0–1 keV, below the simulation cut", "empty"],
    ["h_hit_E_keV_map", "hits per crossing per keV, by layer (not a spatial map)", "labels missing"],
    ["h_hit_particle_ID, h_hit_particle_ID_map", "hits per crossing by species of the linked stored MC particle (the nearest stored ancestor)", "labels missing"],
    ["h_hit_particle_pt, h_hit_particle_E, h_hit_particle_ID_E_MeV", "properties at production of that linked particle, in GeV with 0.1 GeV bins (97–99.8% in the first bin); not the particle at the sensor", "misleading"],
    ["h_hit_particle_eta", "direction (η) at production of that linked particle, hits per crossing per 0.1; not the direction at the sensor (median |η| 1.7 for L3 hits, 3.2 for L4 hits)", "mislabelled"],
    ["hist_hit_t", "hits per crossing per 0.1 ns, 0–20 ns; 16% (L3) and 23% (L4) of the hits are beyond the axis", "labels missing"],
    ["hist_hit_t_corr, hist_hit_t_corr_map", "hit time minus light time from the origin; negative for 47% of the hits because the pairs do not start at the origin; the map version starts at 0 and loses them", "labels missing / misleading"],
    ["h_hit_x_mm, h_hit_y_mm, h_hit_z_mm, h_hit_r_mm, h_hit_eta", "hits per crossing per bin (0.11 mm, 0.086 mm, 0.05); “r” is signed by x", "unit missing"],
    ["h_hit_rateDensity_layer0–4_VS_eta", "hits/BX/cm² of an ideal cylinder × 52 MHz × 3 (safety) × 5 (cluster size): 19–23× the simulated hit rate at 37.04 MHz", "misleading"],
    ["hist_z_density_vs_layer0–4_mm, hist_phi_density_vs_layer0–4", "hits per crossing per 8.64 mm (z) or 0.0707 rad (φ) bin; not divided by area", "labels missing"],
]
AUD_NOTE = ('"hits" = SimTrackerHits; hit-level histograms are filled with weight 1/19,673 per hit, so their contents are '
            "hits per bunch crossing per bin; the occupancy, firing-cell, rate-density and pile-up families are per-crossing "
            "or specially weighted quantities, as listed. Relabelled versions: outputs/BIB/bibstudies_LCC_labelled/<name>.png "
            "and the web copy under truth_ntuple/bibstudies_labelled/.")
for k, chunk in enumerate((AUDIT[:12], AUDIT[12:])):
    s = new_slide(f"Backup: axis audit of all 69 bib-studies histograms ({k + 1}/2)", notes=AUD_NOTE)
    table(s, 0.62, 1.2, 12.1, [0.27, 0.58, 0.15], chunk,
          ["histogram (VertexBarrelCollection)", "y (or colour) axis really is  (“hits” = SimTrackerHits per bunch crossing; relabelled plots in truth_ntuple/bibstudies_labelled/)", "verdict"],
          size=11, row_h=0.4, hrow_h=0.4, name=f"Audit table {k + 1}")
    footer(s)

import datetime, zipfile, shutil
cp = prs.core_properties
cp.title = "Beam-induced background for the MAPS vertex demonstrator: answers to the 4 Aug comments"
cp.subject = "IPC @ Z, LCC V106.2, central mdi_v2 simulation"
cp.author = cp.last_modified_by = "Ang Li"
cp.created = cp.modified = datetime.datetime(2026, 10, 5, 22, 0, 0)
cp.revision = 1
prs.save(str(OUT))
# refresh docProps/app.xml (python-pptx keeps the template's slide count and titles)
titles = [sh.shapes.title.text_frame.text if sh.shapes.title is not None else "" for sh in prs.slides]
esc = lambda t: t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
app = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
       '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
       'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
       '<Template>BNL</Template><TotalTime>0</TotalTime><Application>Microsoft Office PowerPoint</Application>'
       '<PresentationFormat>Widescreen</PresentationFormat>'
       f'<Slides>{len(titles)}</Slides><Notes>{len(titles)}</Notes><HiddenSlides>0</HiddenSlides><MMClips>0</MMClips><ScaleCrop>false</ScaleCrop>'
       '<HeadingPairs><vt:vector size="4" baseType="variant"><vt:variant><vt:lpstr>Theme</vt:lpstr></vt:variant><vt:variant><vt:i4>1</vt:i4></vt:variant>'
       f'<vt:variant><vt:lpstr>Slide Titles</vt:lpstr></vt:variant><vt:variant><vt:i4>{len(titles)}</vt:i4></vt:variant></vt:vector></HeadingPairs>'
       f'<TitlesOfParts><vt:vector size="{len(titles) + 1}" baseType="lpstr"><vt:lpstr>BNL</vt:lpstr>'
       + "".join(f"<vt:lpstr>{esc(t)}</vt:lpstr>" for t in titles) +
       '</vt:vector></TitlesOfParts><Company></Company><LinksUpToDate>false</LinksUpToDate><SharedDoc>false</SharedDoc>'
       '<HyperlinksChanged>false</HyperlinksChanged><AppVersion>16.0000</AppVersion></Properties>')
tmp = OUT.with_suffix(".tmp.pptx")
with zipfile.ZipFile(OUT) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = app.encode("utf-8") if item.filename == "docProps/app.xml" else zin.read(item.filename)
        zout.writestr(item, data)
shutil.move(str(tmp), str(OUT))
print("wrote", OUT, "slides", len(prs.slides))
